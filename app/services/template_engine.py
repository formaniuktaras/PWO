from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.entities import DocType, Document, Event, EventItem, EventUnit, Nomenclature, Unit, User
from app.services.audit_service import AuditService
from app.services.storage_service import StorageService


class TemplateEngine:
    def __init__(self, session: Session, storage: StorageService, templates_root: str, user: User):
        self.session = session
        self.storage = storage
        self.templates_root = Path(templates_root)
        self.audit = AuditService(session, user)

    def generate(self, doc_type: DocType, event_id: int, context: dict[str, str]) -> Path:
        event = self.session.get(Event, event_id)
        if not event:
            raise ValueError("Event not found")
        primary = self.session.query(Unit).join(EventUnit, EventUnit.unit_id == Unit.id).filter(
            EventUnit.event_id == event_id, EventUnit.is_primary.is_(True)
        ).first()
        primary_code = primary.code if primary else "MAIN"

        template_file = Path(doc_type.template_path)
        if not template_file.is_absolute():
            template_file = self.templates_root / doc_type.template_path
        if not template_file.exists():
            raise FileNotFoundError(f"Template missing: {template_file}")

        reg_date = context.get("REG_DATE", event.event_date.isoformat())
        event_base = self.storage.ensure_event_dirs(event_id)
        filename = self.storage.build_doc_name(
            doc_date=event.event_date,
            primary_unit_code=primary_code,
            doc_type=doc_type.code,
            doc_no=context.get("DOC_NO", "draft"),
            reg_date=event.event_date.fromisoformat(reg_date) if isinstance(reg_date, str) else event.event_date,
            ext=doc_type.extension,
        )
        out_file = event_base / "documents" / filename
        shutil.copy2(template_file, out_file)

        ext = doc_type.extension.lower()
        if ext == "docx":
            self._fill_docx(out_file, event_id, context)
        elif ext in {"xlsx", "xlsm"}:
            self._fill_xlsx(out_file, context)
        # docm copied as-is in MVP to preserve macros

        rel = self.storage.relative_to_root(out_file)
        doc = Document(
            event_id=event_id,
            doc_type_id=doc_type.id,
            doc_no=context.get("DOC_NO"),
            doc_date=event.event_date,
            reg_date=event.event_date,
            file_path=rel,
            sha256=self.storage.hash(out_file),
        )
        self.session.add(doc)
        self.session.flush()
        self.audit.log("generate", "documents", str(doc.id), f"Generated from template {template_file.name}")
        return out_file

    def _fill_docx(self, path: Path, event_id: int, context: dict[str, str]) -> None:
        d = DocxDocument(path)
        token = "{{ITEM_TABLE}}"
        for p in list(d.paragraphs):
            for key, val in context.items():
                p.text = p.text.replace(f"{{{{{key}}}}}", str(val))
            if token in p.text:
                p.text = p.text.replace(token, "")
                table = d.add_table(rows=1, cols=4)
                table.rows[0].cells[0].text = "ID"
                table.rows[0].cells[1].text = "Kind"
                table.rows[0].cells[2].text = "Item"
                table.rows[0].cells[3].text = "Qty"
                items = self.session.scalars(select(EventItem).where(EventItem.event_id == event_id)).all()
                for item in items:
                    row = table.add_row().cells
                    row[0].text = str(item.id)
                    row[1].text = item.kind.value
                    row[2].text = str(item.object_id or item.nom_id or "")
                    row[3].text = str(item.qty)
        d.save(path)

    def _fill_xlsx(self, path: Path, context: dict[str, str]) -> None:
        wb = load_workbook(path, keep_vba=True)
        for n, defn in wb.defined_names.items():
            if n in context:
                for sheet_name, cell in defn.destinations:
                    ws = wb[sheet_name]
                    ws[cell] = context[n]
        wb.save(path)
