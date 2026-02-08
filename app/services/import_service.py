from __future__ import annotations

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.entities import DocType, EngineType, Nomenclature, Service, Unit


class ImportService:
    MODELS = {
        "units": Unit,
        "services": Service,
        "nomenclature": Nomenclature,
        "doc_types": DocType,
    }

    def __init__(self, session: Session):
        self.session = session

    def _rows(self, file_path: str):
        wb = load_workbook(file_path, read_only=True, keep_vba=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(x).strip() if x else "" for x in rows[0]]
        return headers, rows[1:]

    def preview(self, file_path: str, limit: int = 20) -> list[dict]:
        headers, rows = self._rows(file_path)
        return [{headers[i]: row[i] for i in range(len(headers))} for row in rows[:limit]]

    def suggest_mapping(self, file_path: str, target: str) -> dict[str, str]:
        headers, _ = self._rows(file_path)
        model = self.MODELS[target]
        valid = set(model.__table__.columns.keys()) | {"parent_code", "parent_name"}
        mapping = {}
        for h in headers:
            key = h.strip().lower()
            if key in valid:
                mapping[h] = key
        return mapping

    def import_rows(self, file_path: str, target: str, mapping: dict[str, str]) -> int:
        model = self.MODELS[target]
        headers, rows = self._rows(file_path)
        count = 0
        units_parent_buffer: list[tuple[Unit, str, str]] = []
        for row in rows:
            payload = {dest: row[headers.index(src)] for src, dest in mapping.items() if src in headers and dest in model.__table__.columns}
            payload = {k: v for k, v in payload.items() if v is not None}
            code = payload.get("code")
            name = payload.get("name")
            lookup_col = model.code if hasattr(model, "code") and code else model.name
            lookup_val = str(code) if code else str(name) if name else None
            if not lookup_val:
                continue
            existing = self.session.scalar(select(model).where(lookup_col == lookup_val))
            if target == "doc_types":
                if "engine_type" in payload and payload["engine_type"]:
                    payload["engine_type"] = EngineType(str(payload["engine_type"]))
            if existing:
                for k, v in payload.items():
                    setattr(existing, k, v)
                obj = existing
            else:
                obj = model(**payload)
                self.session.add(obj)
                self.session.flush()
            if target == "units":
                parent_code_src = next((src for src, dest in mapping.items() if dest == "parent_code"), None)
                parent_name_src = next((src for src, dest in mapping.items() if dest == "parent_name"), None)
                p_code = str(row[headers.index(parent_code_src)]).strip() if parent_code_src and row[headers.index(parent_code_src)] else ""
                p_name = str(row[headers.index(parent_name_src)]).strip() if parent_name_src and row[headers.index(parent_name_src)] else ""
                if p_code or p_name:
                    units_parent_buffer.append((obj, p_code, p_name))
            count += 1

        for unit, parent_code, parent_name in units_parent_buffer:
            parent = None
            if parent_code:
                parent = self.session.scalar(select(Unit).where(Unit.code == parent_code))
            if not parent and parent_name:
                parent = self.session.scalar(select(Unit).where(Unit.short_name == parent_name))
            if parent:
                unit.parent_id = parent.id
        return count
